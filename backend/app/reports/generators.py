import io
import json
import logging
from typing import Dict, Any, List, Optional
from datetime import datetime, date, timedelta
import pandas as pd
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

logger = logging.getLogger("backend.reports.generators")

# -------------------------------------------------------------
# Multi-Format Exporters
# -------------------------------------------------------------

def export_to_csv(df: pd.DataFrame) -> bytes:
    """Exports a pandas DataFrame to CSV bytes."""
    output = io.StringIO()
    df.to_csv(output, index=False)
    return output.getvalue().encode("utf-8")

def export_to_json(data: Dict[str, Any]) -> bytes:
    """Exports structured data to formatted JSON bytes."""
    # Convert numpy types to python native
    def default_serializer(o):
        if isinstance(o, (np.integer, int)):
            return int(o)
        if isinstance(o, (np.floating, float)):
            return float(o)
        if isinstance(o, (np.ndarray, list)):
            return list(o)
        if isinstance(o, (datetime, date)):
            return o.isoformat()
        return str(o)

    json_str = json.dumps(data, indent=2, default=default_serializer)
    return json_str.encode("utf-8")

def export_to_xlsx(sheets: Dict[str, pd.DataFrame], title: str, metadata: Optional[Dict[str, Any]] = None) -> bytes:
    """
    Creates an Excel workbook (.xlsx) with professional enterprise styling,
    custom headers, auto-fit column widths, and optional summary worksheet.
    """
    wb = Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Styles
    header_fill = PatternFill(start_color="0D5C3A", end_color="0D5C3A", fill_type="solid") # Dark Forest Green
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    meta_title_font = Font(name="Calibri", size=14, bold=True, color="0D5C3A")
    bold_font = Font(name="Calibri", size=11, bold=True)
    regular_font = Font(name="Calibri", size=11)
    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0")
    )

    # Sheet 1: Summary / Metadata (if provided)
    if metadata:
        ws_meta = wb.create_sheet(title="Executive Summary")
        ws_meta.views.sheetView[0].showGridLines = True
        ws_meta.append([f"{title} - Executive Summary"])
        ws_meta["A1"].font = meta_title_font
        ws_meta.append([])
        ws_meta.append(["Parameter", "Value"])
        ws_meta["A3"].font = header_font
        ws_meta["A3"].fill = header_fill
        ws_meta["B3"].font = header_font
        ws_meta["B3"].fill = header_fill

        row_idx = 4
        for k, v in metadata.items():
            readable_key = k.replace("_", " ").title()
            val_str = str(v)
            ws_meta.append([readable_key, val_str])
            ws_meta[f"A{row_idx}"].font = bold_font
            ws_meta[f"B{row_idx}"].font = regular_font
            ws_meta[f"A{row_idx}"].border = thin_border
            ws_meta[f"B{row_idx}"].border = thin_border
            row_idx += 1

        ws_meta.column_dimensions["A"].width = 32
        ws_meta.column_dimensions["B"].width = 40

    # Data Sheets
    for sheet_name, df in sheets.items():
        # Truncate sheet name to 31 chars max (Excel limitation)
        safe_name = sheet_name[:31]
        ws = wb.create_sheet(title=safe_name)
        ws.views.sheetView[0].showGridLines = True

        # Header Row
        headers = list(df.columns)
        ws.append(headers)
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border

        # Data Rows
        for row_data in df.itertuples(index=False):
            ws.append(list(row_data))

        # Formatting and Auto Column Sizing
        for col_idx, col in enumerate(df.columns, 1):
            col_letter = get_column_letter(col_idx)
            max_len = max(len(str(col)), max((len(str(val)) for val in df[col].dropna().head(100)), default=0))
            ws.column_dimensions[col_letter].width = max(12, min(max_len + 4, 45))

            # Apply cell borders & number formatting
            for row_idx in range(2, len(df) + 2):
                c = ws.cell(row=row_idx, column=col_idx)
                c.font = regular_font
                c.border = thin_border
                val = c.value
                if isinstance(val, (int, np.integer)):
                    c.number_format = "#,##0"
                elif isinstance(val, (float, np.floating)):
                    c.number_format = "#,##0.00"

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()

def export_to_html_report(
    title: str,
    subtitle: str,
    metadata: Dict[str, Any],
    tables: List[Dict[str, Any]]
) -> bytes:
    """
    Generates a print-optimized, responsive HTML document.
    Suitable for browser viewing or saving/printing to PDF.
    """
    meta_rows_html = ""
    for k, v in metadata.items():
        readable_k = k.replace("_", " ").title()
        meta_rows_html += f"""
        <div class="flex justify-between py-1.5 border-b border-gray-100 text-sm">
            <span class="text-gray-500 font-medium">{readable_k}:</span>
            <span class="text-gray-900 font-semibold">{v}</span>
        </div>
        """

    tables_html = ""
    for t in tables:
        t_title = t.get("title", "")
        t_desc = t.get("description", "")
        df: pd.DataFrame = t.get("dataframe")

        if df is not None and not df.empty:
            headers = "".join(f"<th class='px-3 py-2 text-left text-xs font-semibold text-gray-700 bg-gray-50 uppercase tracking-wider border-b'>{col}</th>" for col in df.columns)
            
            rows_html = ""
            for _, row in df.head(100).iterrows(): # Render up to 100 rows in HTML preview
                cols_html = ""
                for val in row:
                    val_str = str(val)
                    # Pill styling for bands
                    if "within_10" in val_str or "EXCELLENT" in val_str:
                        content = f"<span class='px-2 py-0.5 text-xs font-bold rounded-full bg-emerald-100 text-emerald-800'>{val_str}</span>"
                    elif "between_10" in val_str or "COMPLIANT" in val_str:
                        content = f"<span class='px-2 py-0.5 text-xs font-bold rounded-full bg-amber-100 text-amber-800'>{val_str}</span>"
                    elif "beyond_15" in val_str or "HIGH_PENALTY" in val_str:
                        content = f"<span class='px-2 py-0.5 text-xs font-bold rounded-full bg-rose-100 text-rose-800'>{val_str}</span>"
                    else:
                        content = val_str
                    cols_html += f"<td class='px-3 py-2 text-xs text-gray-800 border-b border-gray-100 whitespace-nowrap'>{content}</td>"
                rows_html += f"<tr class='hover:bg-gray-50 transition'>{cols_html}</tr>"

            tables_html += f"""
            <div class="mt-8">
                <h3 class="text-lg font-bold text-gray-900 mb-1">{t_title}</h3>
                {f"<p class='text-xs text-gray-500 mb-3'>{t_desc}</p>" if t_desc else ""}
                <div class="overflow-x-auto shadow-sm border border-gray-200 rounded-lg">
                    <table class="min-w-full divide-y divide-gray-200">
                        <thead><tr>{headers}</tr></thead>
                        <tbody class="bg-white divide-y divide-gray-100">{rows_html}</tbody>
                    </table>
                </div>
            </div>
            """

    generated_time = datetime.now().strftime("%Y-%m-%d %H:%M:%S IST")

    html_content = f"""<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>{title}</title>
    <script src="https://cdn.tailwindcss.com"></script>
    <style>
        @media print {{
            body {{ -webkit-print-color-adjust: exact; print-color-adjust: exact; font-size: 11px; }}
            .no-print {{ display: none; }}
            .page-break {{ page-break-after: always; }}
        }}
    </style>
</head>
<body class="bg-gray-50 text-gray-900 font-sans p-6 md:p-10">
    <div class="max-w-6xl mx-auto bg-white rounded-xl shadow-md p-8 border border-gray-200">
        <!-- Header -->
        <div class="flex flex-col md:flex-row justify-between items-start md:items-center pb-6 border-b border-gray-200 gap-4">
            <div>
                <div class="flex items-center gap-2">
                    <span class="inline-block w-3 h-3 bg-emerald-600 rounded-full"></span>
                    <span class="text-xs font-bold uppercase tracking-widest text-emerald-800">Govt / Regulatory Compliance Portal</span>
                </div>
                <h1 class="text-2xl md:text-3xl font-black text-gray-900 mt-1">{title}</h1>
                <p class="text-sm text-gray-600 mt-0.5">{subtitle}</p>
            </div>
            <div class="text-right no-print">
                <button onclick="window.print()" class="px-4 py-2 bg-emerald-700 hover:bg-emerald-800 text-white text-xs font-semibold rounded-lg shadow-sm transition">
                    Print to PDF
                </button>
                <div class="text-[11px] text-gray-400 mt-1">Generated: {generated_time}</div>
            </div>
        </div>

        <!-- Summary Grid -->
        <div class="mt-6 bg-slate-50 rounded-lg p-5 border border-slate-200">
            <h2 class="text-sm font-bold uppercase tracking-wider text-slate-700 mb-3">Executive Summary</h2>
            <div class="grid grid-cols-1 md:grid-cols-2 lg:grid-cols-3 gap-x-8 gap-y-1">
                {meta_rows_html}
            </div>
        </div>

        <!-- Tables -->
        {tables_html}

        <!-- Footer -->
        <div class="mt-12 pt-4 border-t border-gray-200 text-center text-xs text-gray-400">
            Generated by AI-Powered Renewable Forecasting Platform • Standard Grid Code & CERC DSM Module
        </div>
    </div>
</body>
</html>"""

    return html_content.encode("utf-8")
